"""Fill the COMP9444 peer group assessment form from the submitter's own text.

Column 1 is the submitter, per the form's instructions. The three 23% members
share one paragraph in the source text; it is written into each of their columns
so every column stands on its own, which is how the form is read.

Fields the submitter must supply by hand are left blank and listed at the end.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt

SRC = Path("form.docx")
DST = Path("Peer_Group_Assessment_Form_-_Yuyi_Zhu.docx")

TEAM = "Austral"
ME = ("Yuyi Zhu", "z5629988")

# column order: submitter first, then the rest
MEMBERS = [
    ("Yuyi Zhu", "z5629988", "26%"),
    ("Ian Zhao", "z5666368", "23%"),
    ("Songning Liu", "z5668661", "23%"),
    ("Maoqin Liu", "z5653393", "23%"),
    ("Yilin", "z5593916", "5%"),
]

CORE = (
    "Formed the core analytical and research team together with Yuyi. "
    "Collaboratively reviewed the initial experimental outputs, identified the "
    "issue of unusually high accuracy, and actively participated in tutorial "
    "discussions to consult with the tutor. Together with Yuyi, researched and "
    "selected the supplementary field dataset, formulated the project improvement "
    "strategy, and co-authored the final report and presentation slides."
)

EXPLANATIONS = [
    ("Acted as the primary technical developer while fully participating in all "
     "team discussions. Built the foundational codebase, executed the experimental "
     "pipeline, and implemented the integration of the supplementary field dataset. "
     "Along with Ian, Songning and Maoqin, actively contributed to analyzing the "
     "experimental results, engaging in tutorial discussions with the tutor, and "
     "preparing the final report and presentation."),
    CORE,
    CORE,
    CORE,
    ("Had minimal involvement in the project's core development and implementation "
     "due to absences from classes and a lack of proactive communication. Provided "
     "highly limited input during the initial codebase review and remained largely "
     "inactive until the group initiated contact during the final preparation stages."),
]


def write_cell(cell, text, size=Pt(8), bold=False):
    """Put text in an empty form cell, matching the form's Arial styling.

    The blank template carries spare empty paragraphs in some cells -- writing
    space for a printed form. Left in, they inflate the row and push the table
    onto a second page, so drop the ones after the first.
    """
    para = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.size = size
    run.font.bold = bold


def append_after(paragraph, text):
    """Append a value to a label paragraph, inheriting the label's formatting."""
    src = paragraph.runs[-1]
    run = paragraph.add_run(text)
    run.font.name = src.font.name
    run.font.size = src.font.size
    run.font.bold = False


def main():
    doc = Document(SRC)

    # header block: Team / Mentor / Student name / Student zid
    header = doc.tables[0].rows[0].cells[0]
    header_values = {"Team Name:": f"  {TEAM}",
                     "Student Name:": f"  {ME[0]}",
                     "Student zid:": f"  {ME[1]}"}
    for para in header.paragraphs:
        for label, value in header_values.items():
            if para.text.strip() == label:
                append_after(para, value)

    # member grid
    grid = doc.tables[1]
    for i, (name, zid, pct) in enumerate(MEMBERS, start=1):
        write_cell(grid.rows[0].cells[i], name, Pt(9), bold=True)
        write_cell(grid.rows[1].cells[i], zid, Pt(9))
        write_cell(grid.rows[2].cells[i], pct, Pt(9), bold=True)
        write_cell(grid.rows[3].cells[i], EXPLANATIONS[i - 1], Pt(7.5))

    # declaration
    for para in doc.paragraphs:
        stripped = para.text.strip()
        if stripped.startswith("Name: ---"):
            para.runs[-1].text = f" {ME[0]}"
        elif stripped.startswith("zid: : ---"):
            para.runs[-1].text = f" {ME[1]}"

    doc.save(DST)
    total = sum(int(m[2].rstrip("%")) for m in MEMBERS)
    print(f"wrote {DST}")
    print(f"contribution total: {total}%")
    blanks = [n for n, z, _ in MEMBERS if not z]
    print("\nleft blank for you to complete by hand:")
    print("  - Mentor Name (your tutor)")
    if blanks:
        print(f"  - zID for: {', '.join(blanks)}")
    print("  - Signature")


if __name__ == "__main__":
    main()
